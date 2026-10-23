"""Motor local de conversión. Copyright 2026. SPDX-License-Identifier: AGPL-3.0-or-later."""
from __future__ import annotations

import os
import re
import statistics
import tempfile
import threading
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import pymupdf
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


class Cancelled(Exception):
    pass


@dataclass
class Options:
    page_markers: bool = True
    pdf_tables: bool = True
    pdf_headings: bool = True


@dataclass
class Result:
    source: str
    output: str
    units: int = 0
    empty_pages: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def escape(text: str) -> str:
    return re.sub(r"([\\`*_{}\[\]<>#|])", r"\\\1", text).replace("\x00", "")


def check_cancel(event: threading.Event):
    if event.is_set():
        raise Cancelled("Conversión cancelada")


def table_md(rows):
    if not rows:
        return ""
    width = max(map(len, rows))
    if not width:
        return ""
    lines = []
    for row in rows:
        cells = [(c or "").replace("\r", "").replace("\n", "<br>") for c in row]
        cells += [""] * (width - len(cells))
        lines.append("| " + " | ".join(cells) + " |")
    lines.insert(1, "| " + " | ".join(["---"] * width) + " |")
    return "\n".join(lines) + "\n\n"


def pdf_to_md(source, out, options, result, progress, cancel):
    with pymupdf.open(source) as doc:
        if doc.needs_pass:
            raise ValueError("PDF protegido con contraseña. Guarda primero una copia desbloqueada.")
        if not len(doc):
            raise ValueError("El PDF no contiene páginas.")
        # Only sample a few pages; never retain all page layouts in memory.
        sizes = Counter()
        for i in sorted(set([0, len(doc) // 4, len(doc) // 2, len(doc) - 1])):
            check_cancel(cancel)
            for block in doc[i].get_text("dict", flags=pymupdf.TEXTFLAGS_TEXT)["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        sizes[round(span["size"])] += len(span["text"].strip())
        baseline = sizes.most_common(1)[0][0] if sizes else 11
        for index, page in enumerate(doc):
            check_cancel(cancel)
            progress(index, len(doc), f"Página {index + 1} de {len(doc)}")
            if options.page_markers:
                out.write(f"<!-- Página {index + 1} -->\n\n")
            layout = page.get_text("dict", flags=pymupdf.TEXTFLAGS_TEXT, sort=True)
            text_blocks = [b for b in layout["blocks"] if b.get("type") == 0]
            raw = " ".join(s["text"] for b in text_blocks for l in b["lines"] for s in l["spans"])
            if not raw.strip():
                result.empty_pages.append(index + 1)
                out.write("> Página sin texto extraíble. Puede requerir OCR.\n\n")
            entries = []
            table_boxes = []
            if options.pdf_tables and raw.strip():
                try:
                    found = page.find_tables()
                    for table in found.tables:
                        rows = table.extract()
                        if rows and len(rows) >= 2 and len(rows[0]) >= 2:
                            table_boxes.append(pymupdf.Rect(table.bbox))
                            entries.append((table.bbox[1], table.bbox[0], table_md(
                                [[escape(c or "") for c in r] for r in rows])))
                except Exception:
                    if "No se pudieron interpretar algunas tablas PDF; se conserva su texto." not in result.warnings:
                        result.warnings.append("No se pudieron interpretar algunas tablas PDF; se conserva su texto.")
            for block in text_blocks:
                # Exclude individual lines in tables, not a whole block that may cross a table boundary.
                lines = []
                spans = []
                for line in block["lines"]:
                    rect = pymupdf.Rect(line["bbox"])
                    center = pymupdf.Point((rect.x0 + rect.x1) / 2, (rect.y0 + rect.y1) / 2)
                    if any(center in box for box in table_boxes):
                        continue
                    value = "".join(s["text"] for s in line["spans"]).strip()
                    if value:
                        lines.append(value)
                        spans.extend(line["spans"])
                if not lines:
                    continue
                text = " ".join(lines)
                # Preserve a line break for recognizable lists.
                is_list = all(re.match(r"^(?:[•●▪‣-]|\d+[.)])\s", line) for line in lines)
                if is_list:
                    value = "\n".join(re.sub(r"^[•●▪‣]\s*", "- ", line) for line in lines)
                else:
                    value = escape(text)
                    size = statistics.median(s["size"] for s in spans) if spans else baseline
                    if options.pdf_headings and len(text) < 180 and size >= baseline * 1.18:
                        level = 1 if size >= baseline * 1.65 else 2 if size >= baseline * 1.35 else 3
                        value = "#" * level + " " + value
                entries.append((block["bbox"][1], block["bbox"][0], value + "\n\n"))
            for _, _, value in sorted(entries):
                out.write(value)
            result.units += 1
        if result.empty_pages:
            result.warnings.append(f"{len(result.empty_pages)} página(s) sin texto extraíble; revisa si requieren OCR.")
        result.warnings.append("PDF: el orden de lectura, títulos y tablas se infieren; revisa diseños con columnas o tablas complejas. Las imágenes no se exportan.")


def run_md(node):
    parts = []
    for child in node:
        if child.tag == qn("w:t"):
            parts.append(escape(child.text or ""))
        elif child.tag in (qn("w:br"), qn("w:cr")):
            parts.append("  \n")
        elif child.tag == qn("w:tab"):
            parts.append("    ")
        elif child.tag == qn("w:footnoteReference"):
            parts.append(f"[^nota-{child.get(qn('w:id'))}]")
        elif child.tag == qn("w:endnoteReference"):
            parts.append(f"[^final-{child.get(qn('w:id'))}]")
    text = "".join(parts)
    props = node.find(qn("w:rPr"))
    if props is not None and text.strip():
        def enabled(name):
            el = props.find(qn(name))
            return el is not None and el.get(qn("w:val"), "true") not in ("0", "false", "off")
        lead = text[:len(text) - len(text.lstrip())]
        trail = text[len(text.rstrip()):]
        text = text.strip()
        if enabled("w:b"):
            text = "**" + text + "**"
        if enabled("w:i"):
            text = "*" + text + "*"
        text = lead + text + trail
    return text


def inline_md(element, part):
    chunks = []
    for node in element:
        if node.tag == qn("w:r"):
            chunks.append(run_md(node))
        elif node.tag == qn("w:hyperlink"):
            label = inline_md(node, part)
            rel = part.rels.get(node.get(qn("r:id")))
            anchor = node.get(qn("w:anchor"))
            url = rel.target_ref if rel is not None else "#" + anchor if anchor else ""
            chunks.append(f"[{label}](<{url.replace('>', '%3E').replace('<', '%3C')}>)" if url else label)
        elif node.tag not in (qn("w:pPr"), qn("w:rPr"), qn("w:del")):
            chunks.append(inline_md(node, part))
    return "".join(chunks)


def docx_to_md(source, out, options, result, progress, cancel):
    doc = Document(source)
    blocks = list(doc.iter_inner_content())
    counters = {}
    images = False
    numbering = doc.part.numbering_part.element

    def paragraph_md(p):
        text = inline_md(p._p, p.part).strip()
        if not text:
            return ""
        style = p.style
        level = None
        num = None
        # Follow inherited paragraph properties, including headings/list styles.
        for props in [p._p.pPr] + [s.element.pPr for s in style_chain(style)]:
            if props is None:
                continue
            outline = props.find(qn("w:outlineLvl"))
            if level is None and outline is not None:
                level = int(outline.get(qn("w:val"))) + 1
            if num is None and props.numPr is not None:
                num = props.numPr
        name = style.name.lower() if style is not None else ""
        match = re.search(r"(?:heading|título|titulo)\s*(\d)", name)
        if match:
            level = int(match[1])
        if level is not None and level <= 6:
            return "#" * max(1, level) + " " + text + "\n\n"
        if name in ("title", "título", "titulo"):
            return "# " + text + "\n\n"
        if num is not None and num.numId is not None and num.numId.val != 0:
            num_id = num.numId.val
            indent = num.ilvl.val if num.ilvl is not None else 0
            fmt, start = "bullet", 1
            abstract_id = None
            override = None
            for n in numbering.findall(qn("w:num")):
                if int(n.get(qn("w:numId"))) == num_id:
                    abstract_id = n.find(qn("w:abstractNumId")).get(qn("w:val"))
                    for ov in n.findall(qn("w:lvlOverride")):
                        if int(ov.get(qn("w:ilvl"))) == indent:
                            override = ov.find(qn("w:startOverride"))
            for abstract in numbering.findall(qn("w:abstractNum")):
                if abstract.get(qn("w:abstractNumId")) == abstract_id:
                    for lvl in abstract.findall(qn("w:lvl")):
                        if int(lvl.get(qn("w:ilvl"))) == indent:
                            f = lvl.find(qn("w:numFmt"))
                            s = lvl.find(qn("w:start"))
                            fmt = f.get(qn("w:val")) if f is not None else "bullet"
                            start = int(s.get(qn("w:val"))) if s is not None else 1
            if override is not None:
                start = int(override.get(qn("w:val")))
            key = (num_id, indent)
            counters[key] = counters.get(key, start - 1) + 1
            for old in list(counters):
                if old[0] == num_id and old[1] > indent:
                    del counters[old]
            marker = "-" if fmt == "bullet" else f"{counters[key]}."
            return "    " * indent + marker + " " + text + "\n\n"
        if "list bullet" in name:
            return "- " + text + "\n\n"
        if "quote" in name or "cita" in name:
            return "> " + text.replace("\n", "\n> ") + "\n\n"
        return text + "\n\n"

    for i, block in enumerate(blocks):
        check_cancel(cancel)
        if i % 20 == 0:
            progress(i, max(1, len(blocks)), f"Bloque {i + 1} de {len(blocks)}")
        if isinstance(block, Paragraph):
            out.write(paragraph_md(block))
            images = images or bool(block._p.xpath(".//w:drawing | .//w:pict"))
        elif isinstance(block, Table):
            rows = []
            for row in block.rows:
                check_cancel(cancel)
                rows.append(["<br>".join(inline_md(p._p, p.part).strip().replace("\n", "<br>") for p in cell.paragraphs) for cell in row.cells])
            out.write(table_md(rows))
        result.units += 1
    # Footnotes and endnotes, stored outside the main document body.
    from lxml import etree
    for part in doc.part.package.parts:
        for suffix, tag, prefix in [("/footnotes.xml", "footnote", "nota"), ("/endnotes.xml", "endnote", "final")]:
            if str(part.partname).endswith(suffix):
                root = etree.fromstring(part.blob)
                for note in root:
                    check_cancel(cancel)
                    note_id = note.get(qn("w:id"))
                    if note_id is None or int(note_id) <= 0:
                        continue
                    text = "\n\n".join(inline_md(p, part).strip() for p in note.findall(qn("w:p")))
                    out.write(f"[^{prefix}-{note_id}]: " + text.replace("\n", "\n    ") + "\n\n")
    if images:
        result.warnings.append("Word contiene imágenes: se exporta el texto, sin las imágenes.")
    result.warnings.append("Word: tablas combinadas/anidadas, cuadros de texto, ecuaciones y formatos de lista especiales pueden requerir revisión. Encabezados y pies no se exportan.")


def style_chain(style):
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        yield style
        style = style.base_style


def convert(source: Path, output_dir: Path, options=None, progress=None, cancel=None) -> Result:
    source, output_dir = Path(source), Path(output_dir)
    if source.suffix.lower() not in (".pdf", ".docx"):
        raise ValueError("Formato no compatible. Usa PDF o Word .docx; guarda los .doc antiguos como .docx.")
    if not source.is_file():
        raise FileNotFoundError(f"No se encontró el archivo: {source}")
    options = options or Options()
    cancel = cancel or threading.Event()
    progress = progress or (lambda *args: None)
    check_cancel(cancel)
    output_dir.mkdir(parents=True, exist_ok=True)
    result = Result(str(source), "")
    fd, temp = tempfile.mkstemp(prefix=".conversion-", suffix=".tmp", dir=output_dir)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as out:
            if source.suffix.lower() == ".pdf":
                pdf_to_md(source, out, options, result, progress, cancel)
            else:
                docx_to_md(source, out, options, result, progress, cancel)
        check_cancel(cancel)
        # Reserve destination exclusively: never replace an existing Markdown.
        number = 1
        while True:
            stem = source.stem if number == 1 else f"{source.stem} ({number})"
            target = output_dir / (stem + ".md")
            try:
                # Exclusive creation also works on FAT/exFAT portable drives.
                with open(target, "x", encoding="utf-8") as dest:
                    try:
                        import shutil
                        with open(temp, encoding="utf-8") as src:
                            shutil.copyfileobj(src, dest, length=1024 * 1024)
                    except BaseException:
                        dest.close()
                        target.unlink(missing_ok=True)
                        raise
                break
            except FileExistsError:
                number += 1
        result.output = str(target)
        progress(result.units, max(1, result.units), "Completado")
        return result
    finally:
        Path(temp).unlink(missing_ok=True)
