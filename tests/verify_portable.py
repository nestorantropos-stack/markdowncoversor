"""End-to-end checks of the compiled executable with Python paths removed."""
import json
import os
import subprocess
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pymupdf
from docx import Document

base = Path(__file__).resolve().parent / "work" / "portable-check"
base.mkdir(parents=True, exist_ok=True)
exe = Path(__file__).resolve().parents[1] / "portable" / "ConversorMarkdown.exe"
env = dict(os.environ)
for name in ("PYTHONPATH", "PYTHONHOME"):
    env.pop(name, None)
env["TEMP"] = env["TMP"] = str(base)
started = time.monotonic()
subprocess.run([str(exe), "--check-ui", str(base / "ui.txt")], env=env, check=True, timeout=45)
assert (base / "ui.txt").read_text(encoding="utf-8") == "UI OK"
print("Standalone GUI: OK", flush=True)
pdf = pymupdf.open()
for i in range(300):
    page = pdf.new_page()
    page.insert_text((50, 60), f"Capitulo {i + 1}", fontsize=22)
    page.insert_text((50, 100), f"Contenido de pagina {i + 1} " * 2, fontsize=11)
pdf.save(base / "libro.pdf")
pdf.close()
doc = Document()
doc.add_heading("Título de prueba", 1)
doc.add_paragraph("Texto local sin instalaciones.")
doc.save(base / "libro.docx")
subprocess.run([str(exe), "--convert", str(base / "libro.pdf"), str(base / "libro.docx"), "--output", str(base / "salida"), "--report", str(base / "report.json")], env=env, check=True, timeout=120)
report = json.loads((base / "report.json").read_text(encoding="utf-8"))
assert len(report) == 2 and all(item["status"] == "ok" for item in report), report
assert report[0]["units"] == 300
assert report[0]["empty_pages"] == []
pdf_text = Path(report[0]["output"]).read_text(encoding="utf-8")
assert pdf_text.count("<!-- Página") == 300
assert "Capitulo 300" in pdf_text
assert "# Título de prueba" in Path(report[1]["output"]).read_text(encoding="utf-8")
assert report[0]["output"] != report[1]["output"]
print(f"Standalone PDF 300 pages + Word: OK ({time.monotonic() - started:.1f} s)", flush=True)
