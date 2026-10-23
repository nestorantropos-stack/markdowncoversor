import json
import sys
import tempfile
import threading
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from converter import Cancelled, Options, convert


class ConversionTests(unittest.TestCase):
    def setUp(self):
        work = Path(__file__).parent / "work"
        work.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=work)
        self.base = Path(self.temp.name)
        self.out = self.base / "salida"

    def tearDown(self):
        self.temp.cleanup()

    def pdf(self, name="ejemplo.pdf", pages=1):
        path = self.base / name
        doc = pymupdf.open()
        for i in range(pages):
            page = doc.new_page()
            page.insert_text((50, 60), f"Capitulo {i + 1}", fontsize=22)
            page.insert_text((50, 100), f"Texto completo de la pagina {i + 1}. Contenido de prueba: conversion.", fontsize=11)
            page.insert_text((50, 125), "\n".join(f"Parrafo {n + 1}: contenido extenso de la pagina {i + 1}, lectura y conversion a Markdown." for n in range(43)), fontsize=10)
        doc.save(path)
        doc.close()
        return path

    def test_300_page_pdf(self):
        events = []
        result = convert(self.pdf(pages=300), self.out, progress=lambda *a: events.append(a))
        text = Path(result.output).read_text(encoding="utf-8")
        self.assertEqual(result.units, 300)
        self.assertEqual(text.count("<!-- Página"), 300)
        self.assertIn("# Capitulo 300", text)
        self.assertIn("pagina 300", text)
        self.assertEqual(result.empty_pages, [])
        self.assertGreater(len(events), 299)

    def test_existing_file_never_overwritten_and_same_stems(self):
        path = self.pdf()
        self.out.mkdir()
        existing = self.out / "ejemplo.md"
        existing.write_text("Conservar", encoding="utf-8")
        result = convert(path, self.out)
        self.assertEqual(existing.read_text(encoding="utf-8"), "Conservar")
        self.assertEqual(Path(result.output).name, "ejemplo (2).md")

    def test_cancel_removes_partial_output(self):
        cancel = threading.Event()
        def progress(*args):
            cancel.set()
        with self.assertRaises(Cancelled):
            convert(self.pdf(pages=10), self.out, progress=progress, cancel=cancel)
        self.assertEqual(list(self.out.iterdir()), [])

    def test_scan_detection_and_password(self):
        blank = self.base / "escaneado.pdf"
        doc = pymupdf.open()
        doc.new_page()
        doc.save(blank)
        doc.close()
        result = convert(blank, self.out)
        self.assertEqual(result.empty_pages, [1])
        self.assertIn("OCR", Path(result.output).read_text(encoding="utf-8"))
        source = self.pdf()
        doc = pymupdf.open(source)
        protected = self.base / "protegido.pdf"
        doc.save(protected, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="secret")
        doc.close()
        with self.assertRaisesRegex(ValueError, "contraseña"):
            convert(protected, self.out)

    def test_corrupted_and_legacy_word(self):
        broken = self.base / "roto.docx"
        broken.write_bytes(b"invalid")
        with self.assertRaises(Exception):
            convert(broken, self.out)
        self.assertEqual(list(self.out.iterdir()), [])
        with self.assertRaisesRegex(ValueError, "docx"):
            convert(self.base / "old.doc", self.out)

    def test_word_structure_and_300_pages(self):
        doc = Document()
        doc.add_heading("Título español", 0)
        doc.add_heading("Sección principal", 1)
        p = doc.add_paragraph()
        p.add_run("Negrita").bold = True
        p.add_run(" y ")
        p.add_run("cursiva").italic = True
        doc.add_paragraph("Uno", style="List Number")
        doc.add_paragraph("Dos", style="List Number")
        doc.add_paragraph("Viñeta", style="List Bullet")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Nombre"
        table.cell(0, 1).text = "Valor"
        table.cell(1, 0).text = "A | B"
        table.cell(1, 1).text = "123"
        p = doc.add_paragraph()
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), doc.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True))
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Enlace"
        run.append(text)
        link.append(run)
        p._p.append(link)
        for i in range(300):
            doc.add_heading(f"Capítulo {i + 1}", 2)
            doc.add_paragraph(f"Párrafo de página {i + 1}. " * 30)
            if i < 299:
                doc.add_page_break()
        path = self.base / "libro.docx"
        doc.save(path)
        result = convert(path, self.out)
        md = Path(result.output).read_text(encoding="utf-8")
        for expected in ["# Título español", "# Sección principal", "**Negrita** y *cursiva*", "1. Uno", "2. Dos", "- Viñeta", "| A \\| B | 123 |", "[Enlace](<https://example.com>)", "## Capítulo 300", "Párrafo de página 300"]:
            self.assertIn(expected, md)

    def test_pdf_table_and_no_duplicates(self):
        doc = pymupdf.open()
        p = doc.new_page()
        for x in (50, 200, 350):
            p.draw_line((x, 100), (x, 190))
        for y in (100, 130, 160, 190):
            p.draw_line((50, y), (350, y))
        for x, y, text in [(60, 120, "Nombre"), (210, 120, "Valor"), (60, 150, "Unico-A"), (210, 150, "10"), (60, 180, "Unico-B"), (210, 180, "20")]:
            p.insert_text((x, y), text)
        path = self.base / "tabla.pdf"
        doc.save(path)
        doc.close()
        result = convert(path, self.out)
        md = Path(result.output).read_text(encoding="utf-8")
        self.assertIn("| Nombre | Valor |", md)
        self.assertEqual(md.count("Unico-A"), 1)

    def test_options(self):
        result = convert(self.pdf(), self.out, Options(False, False, False))
        md = Path(result.output).read_text(encoding="utf-8")
        self.assertNotIn("<!--", md)
        self.assertNotIn("# Capitulo", md)

    def test_footnotes(self):
        doc = Document()
        p = doc.add_paragraph("Texto con nota")
        run = OxmlElement("w:r")
        ref = OxmlElement("w:footnoteReference")
        ref.set(qn("w:id"), "1")
        run.append(ref)
        p._p.append(run)
        original = self.base / "base.docx"
        doc.save(original)
        target = self.base / "notas.docx"
        with zipfile.ZipFile(original) as src, zipfile.ZipFile(target, "w") as dst:
            for entry in src.infolist():
                data = src.read(entry.filename)
                if entry.filename == "[Content_Types].xml":
                    data = data.replace(b"</Types>", b'<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
                if entry.filename == "word/_rels/document.xml.rels":
                    data = data.replace(b"</Relationships>", b'<Relationship Id="rIdNote" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
                dst.writestr(entry, data)
            dst.writestr("word/footnotes.xml", '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>Contenido de la nota.</w:t></w:r></w:p></w:footnote></w:footnotes>')
        result = convert(target, self.out)
        text = Path(result.output).read_text(encoding="utf-8")
        self.assertIn("Texto con nota[^nota-1]", text)
        self.assertIn("[^nota-1]: Contenido de la nota.", text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
