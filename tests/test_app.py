import sys
import tempfile
import time
import tkinter as tk
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import App
from docx import Document


class AppTests(unittest.TestCase):
    def test_gui_batch_recovers_after_error_and_saves_result(self):
        work = Path(__file__).parent / "work"
        work.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=work) as directory:
            base = Path(directory)
            bad = base / "roto.pdf"
            bad.write_bytes(b"invalid")
            good = base / "correcto.docx"
            doc = Document()
            doc.add_heading("Documento correcto", 1)
            doc.add_paragraph("Texto para comprobar el flujo de la interfaz.")
            doc.save(good)
            root = tk.Tk()
            try:
                app = App(root)
                root.update()
                # Check the entire layout is visible at the default window size.
                self.assertLessEqual(app.log.winfo_rooty() + app.log.winfo_height(), root.winfo_rooty() + root.winfo_height())
                with patch("app.filedialog.askopenfilenames", return_value=(str(bad), str(good))):
                    app.add_files()
                app.output.set(str(base / "salida"))
                app.start_conversion()
                deadline = time.monotonic() + 30
                while app.running and time.monotonic() < deadline:
                    root.update()
                    time.sleep(.02)
                self.assertFalse(app.running)
                self.assertEqual(app.done_count, 1)
                self.assertEqual(app.error_count, 1)
                self.assertEqual(app.tree.set(str(bad), "state"), "Error")
                self.assertIn("# Documento correcto", (base / "salida/correcto.md").read_text(encoding="utf-8"))
                self.assertEqual(str(app.start["state"]), "normal")
            finally:
                root.destroy()


if __name__ == "__main__":
    unittest.main(verbosity=2)
